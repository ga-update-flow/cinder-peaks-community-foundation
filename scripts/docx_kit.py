"""Branded python-docx building blocks shared by every generated document."""
from docx import Document
from docx.enum.section import WD_ORIENT
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

CINDER = RGBColor(0xA8, 0x43, 0x2A)
BASALT = RGBColor(0x1F, 0x2A, 0x2E)
PINE = RGBColor(0x2F, 0x5D, 0x50)
STONE = RGBColor(0x5F, 0x6B, 0x6E)
PINE_HEX = "2F5D50"
SAND_HEX = "F4EEE4"
MIST_HEX = "E8EFEC"
BODY_FONT = "Calibri"      # rendered with metric-compatible Carlito by LibreOffice
HEAD_FONT = "Cambria"      # rendered with metric-compatible Caladea by LibreOffice
BOX = "☐"


def _set_font(style, name, size, color=None, bold=None):
    style.font.name = name
    style.font.size = Pt(size)
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.find(qn("w:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.append(rfonts)
    for attr in ("w:ascii", "w:hAnsi", "w:cs", "w:eastAsia"):
        rfonts.set(qn(attr), name)
    if color is not None:
        style.font.color.rgb = color
    if bold is not None:
        style.font.bold = bold


def _field(run, instr):
    for tag, text in (("begin", None), (None, instr), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = text
        run._r.append(el)


def shade(cell, hex_color):
    tcpr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcpr.append(shd)


def _borders(table, color="B9C4C0"):
    tblpr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        el = OxmlElement(f"w:{edge}")
        el.set(qn("w:val"), "single")
        el.set(qn("w:sz"), "4")
        el.set(qn("w:space"), "0")
        el.set(qn("w:color"), color)
        borders.append(el)
    tblpr.append(borders)


class BrandedDoc:
    def __init__(self, ctx, title, doc_type, subtitle=None, landscape=False):
        self.ctx = ctx
        self.site = ctx["site"]
        self.doc = Document()
        d = self.doc
        d.core_properties.title = title
        d.core_properties.author = self.site["org"]["name"]
        d.core_properties.subject = doc_type
        d.core_properties.keywords = "Cinder Peaks Community Foundation; grants; northern Arizona"
        d.core_properties.comments = f"Site {self.site['version_label']}"

        sec = d.sections[0]
        if landscape:
            sec.orientation = WD_ORIENT.LANDSCAPE
            sec.page_width, sec.page_height = Inches(11), Inches(8.5)
        else:
            sec.page_width, sec.page_height = Inches(8.5), Inches(11)
        for side in ("left_margin", "right_margin"):
            setattr(sec, side, Inches(0.9))
        sec.top_margin, sec.bottom_margin = Inches(0.85), Inches(0.8)
        sec.header_distance = sec.footer_distance = Inches(0.4)

        styles = d.styles
        _set_font(styles["Normal"], BODY_FONT, 10.5, BASALT)
        styles["Normal"].paragraph_format.space_after = Pt(6)
        styles["Normal"].paragraph_format.line_spacing = 1.15
        _set_font(styles["Title"], HEAD_FONT, 24, BASALT, True)
        _set_font(styles["Heading 1"], HEAD_FONT, 16, CINDER, True)
        _set_font(styles["Heading 2"], HEAD_FONT, 12.5, PINE, True)
        _set_font(styles["Heading 3"], BODY_FONT, 11, BASALT, True)
        for name, before in (("Heading 1", 16), ("Heading 2", 12), ("Heading 3", 8)):
            styles[name].paragraph_format.space_before = Pt(before)
            styles[name].paragraph_format.space_after = Pt(4)
            styles[name].paragraph_format.keep_with_next = True
        for name in ("List Bullet", "List Number"):
            _set_font(styles[name], BODY_FONT, 10.5, BASALT)
            styles[name].paragraph_format.space_after = Pt(3)

        self._header(doc_type)
        self._footer()
        self._title_block(title, doc_type, subtitle)

    # ---------- page furniture ----------
    def _header(self, doc_type):
        hdr = self.doc.sections[0].header
        p = hdr.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        logo = self.ctx["logo_png"]
        p.add_run().add_picture(str(logo), height=Inches(0.42))
        r = p.add_run("\t\t" + doc_type.upper())
        r.font.size = Pt(8)
        r.font.color.rgb = STONE
        r.font.bold = True
        line = hdr.add_paragraph()
        pbdr = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        for k, v in (("w:val", "single"), ("w:sz", "8"), ("w:space", "1"), ("w:color", "A8432A")):
            bottom.set(qn(k), v)
        pbdr.append(bottom)
        line._p.get_or_add_pPr().append(pbdr)
        line.paragraph_format.space_after = Pt(0)

    def _footer(self):
        c = self.site["contact"]
        ftr = self.doc.sections[0].footer
        p = ftr.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(f"{self.site['org']['name']}  ·  {c['street']}, {c['city_state_zip']}  ·  {c['phone']}  ·  {c['grants_email']}")
        r.font.size = Pt(7.5)
        r.font.color.rgb = STONE
        p2 = ftr.add_paragraph()
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p2.add_run("Page ")
        r.font.size = Pt(7.5)
        r = p2.add_run()
        r.font.size = Pt(7.5)
        _field(r, "PAGE")
        r = p2.add_run(" of ")
        r.font.size = Pt(7.5)
        r = p2.add_run()
        r.font.size = Pt(7.5)
        _field(r, "NUMPAGES")
        r = p2.add_run("   ·   A fictional organization created for website demonstration and testing.")
        r.font.size = Pt(6.5)
        r.font.color.rgb = STONE

    def _title_block(self, title, doc_type, subtitle):
        p = self.doc.add_paragraph()
        r = p.add_run(doc_type.upper())
        r.font.size = Pt(9)
        r.font.bold = True
        r.font.color.rgb = CINDER
        p.paragraph_format.space_after = Pt(2)
        self.doc.add_paragraph(title, style="Title").paragraph_format.space_after = Pt(4)
        if subtitle:
            p = self.doc.add_paragraph()
            r = p.add_run(subtitle)
            r.font.size = Pt(11)
            r.font.color.rgb = STONE
            r.italic = True

    # ---------- content helpers ----------
    def h1(self, text):
        return self.doc.add_heading(text, level=1)

    def h2(self, text):
        return self.doc.add_heading(text, level=2)

    def h3(self, text):
        return self.doc.add_heading(text, level=3)

    def para(self, text="", bold_lead=None, italic=False, size=None, color=None):
        p = self.doc.add_paragraph()
        if bold_lead:
            p.add_run(bold_lead).bold = True
        r = p.add_run(text)
        r.italic = italic
        if size:
            r.font.size = Pt(size)
        if color:
            r.font.color.rgb = color
        return p

    def hint(self, text):
        return self.para(text, italic=True, size=9.5, color=STONE)

    def bullets(self, items, style="List Bullet"):
        for it in items:
            if isinstance(it, tuple):
                p = self.doc.add_paragraph(style=style)
                p.add_run(it[0]).bold = True
                p.add_run(it[1])
            else:
                self.doc.add_paragraph(it, style=style)

    def numbered(self, items):
        # Manual numbering keeps each list restarting at 1 in Word and LibreOffice.
        for i, it in enumerate(items, 1):
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.25)
            p.paragraph_format.space_after = Pt(4)
            if isinstance(it, tuple):
                p.add_run(f"{i}.  {it[0]}. ").bold = True
                p.add_run(it[1])
            else:
                p.add_run(f"{i}.  {it}")

    def checklist(self, items):
        for it in items:
            p = self.doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.first_line_indent = Inches(-0.3)
            p.paragraph_format.space_after = Pt(4)
            r = p.add_run(BOX + "  ")
            r.font.size = Pt(12)
            p.add_run(it)

    def table(self, rows, header=None, widths=None, label_col=False, font_size=9.5, total_row=False):
        cols = len(header or rows[0])
        t = self.doc.add_table(rows=0, cols=cols)
        t.alignment = WD_TABLE_ALIGNMENT.CENTER
        _borders(t)
        if header:
            cells = t.add_row().cells
            for i, h in enumerate(header):
                cells[i].text = ""
                r = cells[i].paragraphs[0].add_run(h)
                r.bold = True
                r.font.size = Pt(font_size)
                r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
                shade(cells[i], PINE_HEX)
            tr = t.rows[0]._tr.get_or_add_trPr()
            th = OxmlElement("w:tblHeader")
            th.set(qn("w:val"), "true")
            tr.append(th)
        for ri, row in enumerate(rows):
            cells = t.add_row().cells
            is_total = total_row and ri == len(rows) - 1
            for i, val in enumerate(row):
                cells[i].text = ""
                para = cells[i].paragraphs[0]
                para.paragraph_format.space_after = Pt(2)
                r = para.add_run(str(val))
                r.font.size = Pt(font_size)
                if (label_col and i == 0) or is_total:
                    r.bold = True
                if label_col and i == 0:
                    shade(cells[i], SAND_HEX)
                elif is_total:
                    shade(cells[i], MIST_HEX)
        if widths:
            for row in t.rows:
                for i, w in enumerate(widths):
                    row.cells[i].width = Inches(w)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)
        return t

    def form(self, fields, widths=(2.2, 4.5), line_height=1):
        """Two-column fill-in form: shaded label, blank answer cell."""
        rows = []
        for f in fields:
            rows.append((f, ""))
        t = self.table(rows, widths=widths, label_col=True)
        if line_height > 1:
            for row in t.rows:
                for _ in range(line_height - 1):
                    row.cells[1].add_paragraph()
        return t

    def answer_box(self, prompt, lines=5, limit=None):
        p = self.doc.add_paragraph()
        p.paragraph_format.keep_with_next = True
        p.add_run(prompt).bold = True
        if limit:
            r = p.add_run(f"  ({limit})")
            r.italic = True
            r.font.color.rgb = STONE
            r.font.size = Pt(9)
        t = self.doc.add_table(rows=1, cols=1)
        _borders(t)
        cell = t.rows[0].cells[0]
        for _ in range(lines - 1):
            cell.add_paragraph()
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def signature_block(self, labels):
        rows = [(l, "") for l in labels]
        self.table(rows, widths=(2.2, 4.5), label_col=True)

    def callout(self, title, text):
        t = self.doc.add_table(rows=1, cols=1)
        cell = t.rows[0].cells[0]
        shade(cell, MIST_HEX)
        p = cell.paragraphs[0]
        r = p.add_run(title)
        r.bold = True
        r.font.color.rgb = PINE
        for line in (text if isinstance(text, list) else [text]):
            cell.add_paragraph(line)
        self.doc.add_paragraph().paragraph_format.space_after = Pt(2)

    def image(self, path, width_in):
        self.doc.add_picture(str(path), width=Inches(width_in))
        self.doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER

    def page_break(self):
        self.doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)

    def contact_block(self, contact, extra=None):
        c = self.site["contact"]
        lines = [
            f"{contact['name']}, {contact['title']}",
            f"Email: {contact['email']}   ·   Phone: {contact['phone']}",
            f"{self.site['org']['name']}, {c['street']}, {c['city_state_zip']}",
        ]
        if extra:
            lines.append(extra)
        self.callout("Questions?", lines)

    def save(self, path):
        path.parent.mkdir(parents=True, exist_ok=True)
        self.doc.save(str(path))
